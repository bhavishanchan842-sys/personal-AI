import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_exact_5_slides(output_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide_layout = prs.slide_layouts[6] # Blank

    NAVY = RGBColor(15, 23, 42)          # Slate 900
    DARK_BLUE = RGBColor(30, 41, 59)     # Slate 800
    EMERALD = RGBColor(16, 185, 129)     # Emerald 500
    TEAL = RGBColor(13, 148, 136)        # Teal 600
    CARD_BG = RGBColor(248, 250, 252)    # Slate 50
    CARD_BORDER = RGBColor(226, 232, 240)# Slate 200
    TEXT_LIGHT = RGBColor(241, 245, 249) # Slate 100
    TEXT_DARK = RGBColor(30, 41, 59)     # Slate 800
    TEXT_BODY = RGBColor(51, 65, 85)     # Slate 700

    def add_header(slide, slide_num, title_text):
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.5), Inches(0.35))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = f"SLIDE {slide_num} OF 5  |  NESTMATE PROJECT REVIEW"
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = TEAL

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.5), Inches(0.7))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(23)
        p_title.font.bold = True
        p_title.font.color.rgb = NAVY

    # =========================================================================
    # SLIDE 1: INTRODUCTION
    # =========================================================================
    s1 = prs.slides.add_slide(slide_layout)
    add_header(s1, "1", "1. Introduction & Background")

    # Project Banner at top of slide 1
    top_banner = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(11.7), Inches(1.15))
    top_banner.fill.solid()
    top_banner.fill.fore_color.rgb = DARK_BLUE
    top_banner.line.fill.background()
    t_tf = top_banner.text_frame
    t_tf.word_wrap = True
    t_tf.margin_left = Inches(0.3)
    t_tf.margin_top = Inches(0.18)
    p = t_tf.paragraphs[0]
    p.text = "Project Title: NestMate — Intelligent Room & PG Finder with Roommate Compatibility"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = EMERALD
    p_sub = t_tf.add_paragraph()
    p_sub.text = "Student Name: [Your Name]   |   Roll No: [Your Roll No]   |   Guide: [Guide / Mam's Name]"
    p_sub.font.size = Pt(12)
    p_sub.font.color.rgb = TEXT_LIGHT

    # Left Column: Problem & Context
    c1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.85), Inches(5.6), Inches(4.2))
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = CARD_BORDER
    tf1 = c1.text_frame
    tf1.word_wrap = True
    tf1.margin_left = Inches(0.3)
    tf1.margin_top = Inches(0.25)
    p = tf1.paragraphs[0]
    p.text = "Background & Motivation"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.space_after = Pt(8)

    for pt in [
        "Rapid Student & Urban Migration: Heavy relocation to major tech and education hubs (Bengaluru, Gurgaon, Pune, Hyderabad, Mumbai).",
        "Financial Reality: High rental prices make shared flats and Paying Guest (PG) accommodations the primary choice.",
        "The Roommate Conflict Problem: Living with mismatched habits causes sleep deprivation, academic stress, lease breakages, and loss of security deposits.",
        "The Core Need: A platform that treats shared housing as a human matching problem, not just a property transaction."
    ]:
        p = tf1.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_BODY
        p.space_after = Pt(6)

    # Right Column: What NestMate Does
    c2 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(2.85), Inches(5.7), Inches(4.2))
    c2.fill.solid()
    c2.fill.fore_color.rgb = CARD_BG
    c2.line.color.rgb = CARD_BORDER
    tf2 = c2.text_frame
    tf2.word_wrap = True
    tf2.margin_left = Inches(0.3)
    tf2.margin_top = Inches(0.25)
    p = tf2.paragraphs[0]
    p.text = "The NestMate Solution"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = TEAL
    p.space_after = Pt(8)

    for pt in [
        "Dual-Integrated Platform: Merges comprehensive PG/room discovery with an algorithmic lifestyle matchmaking engine.",
        "Pre-Booking Transparency: Seekers can view their real-time compatibility match score with current occupants already living in that room before visiting.",
        "6-Axis Lifestyle Profiling: Evaluates sleep cycle, cleanliness, social battery, noise tolerance, work routine, and food habits.",
        "Co-Living Utilities: Includes a Fair Rent Splitter and a 11-Month Digital Roommate Agreement generator."
    ]:
        p = tf2.add_paragraph()
        p.text = "✓ " + pt
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_BODY
        p.space_after = Pt(6)

    # =========================================================================
    # SLIDE 2: LITERATURE SURVEY
    # =========================================================================
    s2 = prs.slides.add_slide(slide_layout)
    add_header(s2, "2", "2. Literature Survey & Related Work")

    table_shape = s2.shapes.add_table(5, 4, Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2))
    t = table_shape.table
    t.columns[0].width = Inches(2.2)
    t.columns[1].width = Inches(3.0)
    t.columns[2].width = Inches(3.2)
    t.columns[3].width = Inches(3.3)

    headers = ["Category / System", "Primary Focus", "Key Strengths", "Identified Limitations & Gaps"]
    for i, h in enumerate(headers):
        cell = t.cell(0, i)
        cell.fill.solid()
        cell.fill.fore_color.rgb = DARK_BLUE
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = TEXT_LIGHT

    survey_data = [
        (
            "Commercial Portals\n(99acres, MagicBricks, Housing)",
            "Real Estate Sales & Whole Apartment Rentals",
            "Extensive property inventory, map search, high traffic, broker networks.",
            "Complete absence of roommate matching. Treats shared accommodation strictly as real estate; no lifestyle assessment."
        ),
        (
            "Managed Co-Living\n(Zolo Stays, Stanza Living)",
            "Managed Hostels & Student Accommodations",
            "Furnished rooms, regular housekeeping, food mess, mobile rent payment.",
            "Zero compatibility screening; arbitrary bed assignments; high deposits and rigid long lock-in contracts."
        ),
        (
            "International Apps\n(Roomi, SpareRoom, Flatmates)",
            "Roommate Pairing in US, UK & Australia",
            "User profiles, budget preference matching, direct messaging.",
            "Lacks Indian context: No PG sharing types (1/2/3 sharing), no food/mess plans, no strict vegetarian dealbreakers."
        ),
        (
            "Academic Research\n(Multi-Attribute Matching)",
            "Theoretical Matchmaking Recommender Systems",
            "Mathematical proof of stable pairing, vector distance algorithms.",
            "Purely theoretical; algorithms exist in papers but are never integrated into an active, functional property booking portal."
        )
    ]
    for row_idx, row in enumerate(survey_data, start=1):
        for col_idx, text in enumerate(row):
            cell = t.cell(row_idx, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(255, 255, 255) if row_idx % 2 == 1 else RGBColor(241, 245, 249)
            p = cell.text_frame.paragraphs[0]
            p.text = text
            p.font.size = Pt(11)
            p.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 3: GAPS IN EXISTING SYSTEMS
    # =========================================================================
    s3 = prs.slides.add_slide(slide_layout)
    add_header(s3, "3", "3. Gaps in Existing Systems")

    gaps = [
        ("Gap 1: Property-Centric vs. Human-Centric", "Existing portals treat accommodation search strictly as real estate transactions. In shared living, tenant satisfaction is 80% determined by roommate interpersonal harmony rather than wall paint or tiles."),
        ("Gap 2: The 'Blind Booking' Risk", "Tenants sign leases without knowing if current occupants are night owls, non-vegetarians, smokers, or have incompatible hygiene habits, causing immediate conflict and lease breaks."),
        ("Gap 3: Missing Indian Context Specifics", "Global platforms do not account for Indian student realities: Boys/Girls/Co-Ed PGs, Food/Mess inclusions (3 meals/day vs self-cooking), Curfew rules, and Zero Brokerage verification."),
        ("Gap 4: Lack of Conflict Prevention Utilities", "No existing platforms provide objective tools like Fair Rent Splitters (adjusting rent for attached bathrooms/balconies) or structured digital co-living agreements.")
    ]
    for idx, (title, desc) in enumerate(gaps):
        gx = Inches(0.8 + (idx % 2) * 5.9)
        gy = Inches(1.6 + (idx // 2) * 2.7)

        gc = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, gx, gy, Inches(5.6), Inches(2.45))
        gc.fill.solid()
        gc.fill.fore_color.rgb = CARD_BG
        gc.line.color.rgb = CARD_BORDER

        gtf = gc.text_frame
        gtf.word_wrap = True
        gtf.margin_left = Inches(0.3)
        gtf.margin_top = Inches(0.25)
        gtf.margin_right = Inches(0.3)

        p = gtf.paragraphs[0]
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = NAVY if idx % 2 == 0 else TEAL
        p.space_after = Pt(6)

        p_desc = gtf.add_paragraph()
        p_desc.text = desc
        p_desc.font.size = Pt(12)
        p_desc.font.color.rgb = TEXT_BODY

    # =========================================================================
    # SLIDE 4: PROBLEM STATEMENT
    # =========================================================================
    s4 = prs.slides.add_slide(slide_layout)
    add_header(s4, "4", "4. Problem Statement")

    # Banner Card
    bcard = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(11.7), Inches(2.4))
    bcard.fill.solid()
    bcard.fill.fore_color.rgb = DARK_BLUE
    bcard.line.fill.background()

    btf = bcard.text_frame
    btf.word_wrap = True
    btf.margin_left = Inches(0.4)
    btf.margin_top = Inches(0.35)
    btf.margin_right = Inches(0.4)

    p = btf.paragraphs[0]
    p.text = "FORMAL PROBLEM STATEMENT"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = EMERALD
    p.space_after = Pt(8)

    p_s = btf.add_paragraph()
    p_s.text = (
        '"To design and develop an intelligent web platform that bridges the gap between '
        'student/PG accommodation discovery and algorithmic roommate lifestyle compatibility, '
        'enabling prospective tenants to evaluate physical property attributes and occupant compatibility '
        'simultaneously before committing to financial deposits or residential agreements."'
    )
    p_s.font.size = Pt(15.5)
    p_s.font.color.rgb = TEXT_LIGHT

    # Bottom Two Cards
    sc1 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.3), Inches(5.6), Inches(2.5))
    sc1.fill.solid()
    sc1.fill.fore_color.rgb = CARD_BG
    sc1.line.color.rgb = CARD_BORDER
    tf = sc1.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.3)
    tf.margin_top = Inches(0.2)
    p = tf.paragraphs[0]
    p.text = "Target Demographics & Stakeholders:"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.space_after = Pt(6)
    for item in [
        "Undergraduate and postgraduate university students.",
        "Entry-level IT and corporate professionals relocating to tech clusters.",
        "Independent PG owners seeking low tenant turnover and peaceful occupancy."
    ]:
        p = tf.add_paragraph()
        p.text = "• " + item
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_BODY
        p.space_after = Pt(4)

    sc2 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(4.3), Inches(5.7), Inches(2.5))
    sc2.fill.solid()
    sc2.fill.fore_color.rgb = CARD_BG
    sc2.line.color.rgb = CARD_BORDER
    tf = sc2.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.3)
    tf.margin_top = Inches(0.2)
    p = tf.paragraphs[0]
    p.text = "Core Challenges Addressed:"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = TEAL
    p.space_after = Pt(6)
    for item in [
        "Eliminating the financial & mental risk of living with incompatible roommates.",
        "Providing complete transparency into current flatmates' daily habits.",
        "Removing broker exploitation with direct, zero-brokerage room search."
    ]:
        p = tf.add_paragraph()
        p.text = "✓ " + item
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_BODY
        p.space_after = Pt(4)

    # =========================================================================
    # SLIDE 5: OBJECTIVES
    # =========================================================================
    s5 = prs.slides.add_slide(slide_layout)
    add_header(s5, "5", "5. Project Objectives")

    oc1 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    oc1.fill.solid()
    oc1.fill.fore_color.rgb = CARD_BG
    oc1.line.color.rgb = CARD_BORDER
    tf = oc1.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.3)
    tf.margin_top = Inches(0.3)
    p = tf.paragraphs[0]
    p.text = "Primary Objectives"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.space_after = Pt(10)

    for obj in [
        "Interactive PG & Room Discovery Engine: Develop filters for Boys/Girls/Co-Ed PGs, single/sharing occupancy, 3 meals included, and budget in ₹ INR.",
        "Algorithmic Compatibility Engine: Implement weighted Euclidean vector distance math across 6 lifestyle dimensions (Sleep, Cleanliness, Social, Work, Noise, Cooking).",
        "Dual-Integration Architecture: Embed real-time occupant compatibility match % directly into property listing cards.",
        "Visual Radar Comparison: Generate 6-axis Radar charts with automatic Green Flags (strengths) and Yellow Flags (habits to discuss)."
    ]:
        p = tf.add_paragraph()
        p.text = "🎯 " + obj
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_BODY
        p.space_after = Pt(8)

    oc2 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    oc2.fill.solid()
    oc2.fill.fore_color.rgb = CARD_BG
    oc2.line.color.rgb = CARD_BORDER
    tf = oc2.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.3)
    tf.margin_top = Inches(0.3)
    p = tf.paragraphs[0]
    p.text = "Secondary / Utility Objectives"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = TEAL
    p.space_after = Pt(10)

    for obj in [
        "Dealbreaker Enforcement: Enforce strict rules for smoking, alcohol, pets, and dietary preferences (Pure Veg vs Non-Veg).",
        "Mathematical Fair Rent Splitter: Build an objective calculator dividing rent based on room size, attached washroom, and balcony.",
        "11-Month Digital Roommate Agreement: Standardize house rules, quiet hours, and chore rotation contracts with print/export options.",
        "High-Performance Architecture: Provide a fast, responsive Single Page Application (SPA) using React 18, Vite, and Python FastAPI."
    ]:
        p = tf.add_paragraph()
        p.text = "⚡ " + obj
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_BODY
        p.space_after = Pt(8)

    prs.save(output_path)
    print(f"[SUCCESS] Exact 5-slide Presentation saved to: {output_path}")

if __name__ == "__main__":
    out_file = r"c:\Drive d\hmmmmm\NestMate_Project_Presentation.pptx"
    create_exact_5_slides(out_file)
