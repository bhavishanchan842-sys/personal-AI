import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation(output_path):
    prs = Presentation()
    # Set 16:9 widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Palette
    NAVY = RGBColor(15, 23, 42)       # Slate 900
    EMERALD = RGBColor(16, 185, 129)  # Emerald 500
    DARK_BLUE = RGBColor(30, 41, 59)  # Slate 800
    TEXT_LIGHT = RGBColor(241, 245, 249) # Slate 100
    TEXT_DARK = RGBColor(30, 41, 59)
    TEXT_MUTED = RGBColor(100, 116, 139)
    ACCENT_BG = RGBColor(248, 250, 252) # Slate 50
    CARD_BG = RGBColor(255, 255, 255)
    BORDER_COLOR = RGBColor(226, 232, 240)
    TEAL = RGBColor(13, 148, 136)

    def add_header(slide, title_text, category_text="NESTMATE PROJECT PRESENTATION"):
        # Category tag
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = TEAL
        
        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.5), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = NAVY

    # ==========================================
    # SLIDE 1: TITLE SLIDE (Dark Theme)
    # ==========================================
    slide_layout = prs.slide_layouts[6] # Blank
    slide1 = prs.slides.add_slide(slide_layout)
    
    # Background rectangle
    bg = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()

    # Title container
    title_box = slide1.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(11), Inches(3.2))
    tf = title_box.text_frame
    tf.word_wrap = True
    
    p_tag = tf.paragraphs[0]
    p_tag.text = "COLLEGE CAPSTONE / ACADEMIC PROJECT"
    p_tag.font.size = Pt(14)
    p_tag.font.bold = True
    p_tag.font.color.rgb = EMERALD
    p_tag.space_after = Pt(14)

    p_title = tf.add_paragraph()
    p_title.text = "NestMate: Intelligent Room & PG Finder"
    p_title.font.size = Pt(36)
    p_title.font.bold = True
    p_title.font.color.rgb = TEXT_LIGHT

    p_sub = tf.add_paragraph()
    p_sub.text = "With Algorithmic Roommate Lifestyle Compatibility Matching"
    p_sub.font.size = Pt(22)
    p_sub.font.color.rgb = RGBColor(148, 163, 184) # Slate 400
    p_sub.space_before = Pt(8)

    # Details Box
    det_box = slide1.shapes.add_textbox(Inches(1.2), Inches(5.2), Inches(10), Inches(1.5))
    tf_det = det_box.text_frame
    p_d1 = tf_det.paragraphs[0]
    p_d1.text = "• Domain: Web Application Development, Information Systems & Algorithmic Matching"
    p_d1.font.size = Pt(14)
    p_d1.font.color.rgb = RGBColor(203, 213, 225)
    
    p_d2 = tf_det.add_paragraph()
    p_d2.text = "• Tech Stack: React.js, Tailwind CSS, Python (FastAPI), SQLite"
    p_d2.font.size = Pt(14)
    p_d2.font.color.rgb = RGBColor(203, 213, 225)
    p_d2.space_before = Pt(6)

    # ==========================================
    # SLIDE 2: INTRODUCTION
    # ==========================================
    slide2 = prs.slides.add_slide(slide_layout)
    add_header(slide2, "Introduction & Background", "Project Overview")

    # Left Card: The Real-World Context
    left_card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    left_card.fill.solid()
    left_card.fill.fore_color.rgb = ACCENT_BG
    left_card.line.color.rgb = BORDER_COLOR
    
    tf_l = left_card.text_frame
    tf_l.word_wrap = True
    tf_l.margin_left = Inches(0.3)
    tf_l.margin_top = Inches(0.3)
    tf_l.margin_right = Inches(0.3)

    p = tf_l.paragraphs[0]
    p.text = "Background & Motivation"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.space_after = Pt(12)

    bullets_left = [
        "Rapid urban migration: Thousands of college students and young professionals relocate to major educational & tech hubs each year.",
        "High cost of individual living makes PG (Paying Guest) and shared flat accommodation the most viable housing option.",
        "Finding an affordable room is only half the struggle — finding compatible people to share living space with is the critical factor for peace of mind and retention.",
        "Frequent roommate conflicts lead to stress, unexpected lease breaks, and financial penalties."
    ]
    for b in bullets_left:
        p = tf_l.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(8)

    # Right Card: What is NestMate?
    right_card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    right_card.fill.solid()
    right_card.fill.fore_color.rgb = ACCENT_BG
    right_card.line.color.rgb = BORDER_COLOR
    
    tf_r = right_card.text_frame
    tf_r.word_wrap = True
    tf_r.margin_left = Inches(0.3)
    tf_r.margin_top = Inches(0.3)
    tf_r.margin_right = Inches(0.3)

    p = tf_r.paragraphs[0]
    p.text = "What is NestMate?"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = TEAL
    p.space_after = Pt(12)

    bullets_right = [
        "Dual-Integrated Platform: A unified web solution combining specialized PG/Room search with an algorithmic compatibility engine.",
        "Pre-Booking Transparency: Enables seekers to view the lifestyle compatibility score with current occupants already residing in a room before visiting.",
        "Multi-Dimensional Profiling: Evaluates 6 key lifestyle axes (Sleep schedule, Cleanliness, Social battery, Noise, Work habits, Meal sharing).",
        "Co-Living Ecosystem: Includes utility tools such as a Fair Rent Splitter and a Digital 11-Month Roommate Agreement generator."
    ]
    for b in bullets_right:
        p = tf_r.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(8)

    # ==========================================
    # SLIDE 3: LITERATURE SURVEY
    # ==========================================
    slide3 = prs.slides.add_slide(slide_layout)
    add_header(slide3, "Literature Survey & Existing Systems Analysis", "Related Work")

    # Table of comparison
    rows, cols = 5, 4
    left = Inches(0.8)
    top = Inches(1.6)
    width = Inches(11.7)
    height = Inches(5.0)

    table_shape = slide3.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table

    table.columns[0].width = Inches(2.2)
    table.columns[1].width = Inches(3.1)
    table.columns[2].width = Inches(3.2)
    table.columns[3].width = Inches(3.2)

    headers = ["Platform / Study", "Focus Area", "Key Features", "Identified Limitations"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.fill.solid()
        cell.fill.fore_color.rgb = DARK_BLUE
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = TEXT_LIGHT

    lit_data = [
        (
            "Commercial Portals\n(e.g., 99acres, MagicBricks)",
            "Real Estate & Property Sales/Rentals",
            "Extensive property inventory, map search, broker directory, price filters.",
            "Complete absence of roommate matching; treats shared accommodation strictly as real estate transactions."
        ),
        (
            "Managed Co-Living\n(e.g., Zolo, Stanza, Nestaway)",
            "Managed PG & Student Housing",
            "Standardized amenities, furnished rooms, mobile app rent payments, housekeeping.",
            "Fixed room allocation with zero tenant-to-tenant compatibility screening; rigid contracts and high lock-in costs."
        ),
        (
            "International Roommate Apps\n(e.g., Roomi, SpareRoom)",
            "Roommate Matching in Western markets",
            "User profiles, budget preference matching, direct messaging.",
            "Not optimized for Indian student context (lacks PG sharing types, meal plans, curfew rules, veg/non-veg dietary dealbreakers)."
        ),
        (
            "Academic Research\n(Multi-Criteria Recommenders)",
            "Stable Matching & Social Matching Models",
            "Gale-Shapley variations, Euclidean distance profiling, personality clustering.",
            "Purely theoretical or isolated algorithms without integration into an operational property search platform."
        )
    ]

    for row_idx, data in enumerate(lit_data, start=1):
        for col_idx, text in enumerate(data):
            cell = table.cell(row_idx, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(255, 255, 255) if row_idx % 2 == 1 else RGBColor(241, 245, 249)
            p = cell.text_frame.paragraphs[0]
            p.text = text
            p.font.size = Pt(11)
            p.font.color.rgb = TEXT_DARK

    # ==========================================
    # SLIDE 4: GAPS IN EXISTING SYSTEMS
    # ==========================================
    slide4 = prs.slides.add_slide(slide_layout)
    add_header(slide4, "Identified Gaps in Existing Systems", "Research & Market Gaps")

    gaps = [
        (
            "1. Disconnect Between Property & People",
            "Existing portals treat accommodation search strictly as a property transaction. However, in shared housing and PGs, tenant satisfaction is 80% determined by co-living interpersonal harmony.",
            TEAL
        ),
        (
            "2. Blind Booking & High Conflict Rates",
            "Prospective tenants are forced to sign leases or book PG beds without knowing the sleep habits, cleanliness standards, or dietary preferences of current occupants, causing frequent post-move disputes.",
            NAVY
        ),
        (
            "3. Lack of Indian PG Context Specifics",
            "Western platforms overlook essential Indian living requirements: Boys/Girls/Co-Ed classifications, food/mess inclusion (3 meals vs self-cooking), strict vegetarian rules, and gate curfew hours.",
            DARK_BLUE
        ),
        (
            "4. Absence of Fair Conflict Prevention Tools",
            "No existing rental portals provide tools to prevent disputes before they occur, such as mathematical rent splitters for unequal room sizes or structured digital co-living agreements.",
            TEAL
        )
    ]

    card_positions = [
        (Inches(0.8), Inches(1.6)),
        (Inches(6.8), Inches(1.6)),
        (Inches(0.8), Inches(4.4)),
        (Inches(6.8), Inches(4.4)),
    ]

    for (title, desc, color), (x, y) in zip(gaps, card_positions):
        card = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.7), Inches(2.4))
        card.fill.solid()
        card.fill.fore_color.rgb = ACCENT_BG
        card.line.color.rgb = BORDER_COLOR
        
        tf_c = card.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = Inches(0.3)
        tf_c.margin_top = Inches(0.25)
        tf_c.margin_right = Inches(0.3)

        p = tf_c.paragraphs[0]
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = color
        p.space_after = Pt(6)

        p_desc = tf_c.add_paragraph()
        p_desc.text = desc
        p_desc.font.size = Pt(12)
        p_desc.font.color.rgb = TEXT_DARK

    # ==========================================
    # SLIDE 5: PROBLEM STATEMENT
    # ==========================================
    slide5 = prs.slides.add_slide(slide_layout)
    add_header(slide5, "Problem Statement", "Core Challenge")

    # Big Banner Card
    banner = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(11.7), Inches(2.3))
    banner.fill.solid()
    banner.fill.fore_color.rgb = DARK_BLUE
    banner.line.fill.background()

    tf_b = banner.text_frame
    tf_b.word_wrap = True
    tf_b.margin_left = Inches(0.4)
    tf_b.margin_top = Inches(0.35)
    tf_b.margin_right = Inches(0.4)

    p_lbl = tf_b.paragraphs[0]
    p_lbl.text = "FORMAL PROBLEM STATEMENT"
    p_lbl.font.size = Pt(12)
    p_lbl.font.bold = True
    p_lbl.font.color.rgb = EMERALD
    p_lbl.space_after = Pt(6)

    p_stmt = tf_b.add_paragraph()
    p_stmt.text = (
        '"To design and develop a comprehensive web platform that bridges the gap between '
        'student/PG accommodation discovery and algorithmic roommate lifestyle compatibility, '
        'enabling seekers to evaluate physical property attributes and tenant compatibility '
        'simultaneously before committing to financial deposits or residential agreements."'
    )
    p_stmt.font.size = Pt(15)
    p_stmt.font.color.rgb = TEXT_LIGHT

    # Two bottom sub-cards: Key Pain Points
    sub1 = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.2), Inches(5.6), Inches(2.6))
    sub1.fill.solid()
    sub1.fill.fore_color.rgb = ACCENT_BG
    sub1.line.color.rgb = BORDER_COLOR
    tf1 = sub1.text_frame
    tf1.word_wrap = True
    tf1.margin_left = Inches(0.3)
    tf1.margin_top = Inches(0.2)

    p = tf1.paragraphs[0]
    p.text = "Primary Challenges Addressed:"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.space_after = Pt(6)

    for pt in [
        "Eliminating the financial risk of renting with incompatible roommates.",
        "Providing transparent visibility into existing occupants' habits.",
        "Reducing broker dependency and excessive brokerage fees."
    ]:
        p = tf1.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(4)

    sub2 = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(4.2), Inches(5.7), Inches(2.6))
    sub2.fill.solid()
    sub2.fill.fore_color.rgb = ACCENT_BG
    sub2.line.color.rgb = BORDER_COLOR
    tf2 = sub2.text_frame
    tf2.word_wrap = True
    tf2.margin_left = Inches(0.3)
    tf2.margin_top = Inches(0.2)

    p = tf2.paragraphs[0]
    p.text = "Target Demographics & Scope:"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = TEAL
    p.space_after = Pt(6)

    for pt in [
        "Undergraduate & postgraduate students shifting to university towns.",
        "Junior IT and corporate professionals migrating to major metro tech clusters.",
        "Independent PG owners seeking low-turnover, satisfied long-term tenants."
    ]:
        p = tf2.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(4)

    # ==========================================
    # SLIDE 6: OBJECTIVES
    # ==========================================
    slide6 = prs.slides.add_slide(slide_layout)
    add_header(slide6, "Project Objectives", "Goals & Deliverables")

    # Left Column: Primary Objectives
    left_o = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    left_o.fill.solid()
    left_o.fill.fore_color.rgb = ACCENT_BG
    left_o.line.color.rgb = BORDER_COLOR
    tfl_o = left_o.text_frame
    tfl_o.word_wrap = True
    tfl_o.margin_left = Inches(0.3)
    tfl_o.margin_top = Inches(0.3)
    tfl_o.margin_right = Inches(0.3)

    p = tfl_o.paragraphs[0]
    p.text = "Primary Objectives"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.space_after = Pt(12)

    primaries = [
        "Develop an Interactive PG & Room Discovery Engine supporting multidimensional filters (Sharing type, Gender specific PGs, Meal plans, Zero brokerage).",
        "Implement a Multi-Attribute Compatibility Matching Engine using weighted vector distance algorithms across 6 core lifestyle dimensions.",
        "Engineer a Dual-Integration Architecture that embeds real-time roommate compatibility scores directly on individual property listing cards.",
        "Provide Visual Conflict Analysis using 6-axis Radar charts, highlighting positive synergies (Green Flags) and points of discussion (Yellow Flags)."
    ]
    for b in primaries:
        p = tfl_o.add_paragraph()
        p.text = "🎯 " + b
        p.font.size = Pt(12.5)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(8)

    # Right Column: Secondary Objectives
    right_o = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    right_o.fill.solid()
    right_o.fill.fore_color.rgb = ACCENT_BG
    right_o.line.color.rgb = BORDER_COLOR
    tfr_o = right_o.text_frame
    tfr_o.word_wrap = True
    tfr_o.margin_left = Inches(0.3)
    tfr_o.margin_top = Inches(0.3)
    tfr_o.margin_right = Inches(0.3)

    p = tfr_o.paragraphs[0]
    p.text = "Secondary / Utility Objectives"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = TEAL
    p.space_after = Pt(12)

    secondaries = [
        "Implement Hard Constraint Dealbreakers (Strict Veg vs Non-Veg, Smoking, Alcohol, Pets) that guarantee mismatch prevention.",
        "Build a Mathematical Fair Rent Splitter based on room square footage, attached washrooms, and private balconies.",
        "Provide a Digital 11-Month Co-Living Agreement generator covering quiet hours, guest policies, and cleaning schedules.",
        "Deliver a modern, fast, responsive Single Page Application (SPA) using React 18, Vite, and FastAPI REST endpoints."
    ]
    for b in secondaries:
        p = tfr_o.add_paragraph()
        p.text = "⚡ " + b
        p.font.size = Pt(12.5)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(8)

    # Save presentation
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    out_file = r"c:\Drive d\hmmmmm\NestMate_Project_Presentation.pptx"
    create_presentation(out_file)
